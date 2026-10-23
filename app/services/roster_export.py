"""The roster as a spreadsheet RD can review, correct and upload back.

One row per client / commission partner: its beneficiary codes split into
writer codes (C…, JN…) and commission codes (CS…, CSJ…), alternate names,
contacts, who can sign in, and the integrity problems found on it
(app/services/integrity.py). The same sheet, edited, is what
app/services/roster_corrections.py reads back. A second sheet lists every
problem; a third explains how to correct.
"""

from __future__ import annotations

import io
from collections import defaultdict
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from sqlalchemy.orm import Session

from app.models.statements import (
    BeneficiaryAccount,
    Contact,
    Writer,
    WriterAlias,
    WriterContact,
    WriterKind,
    WriterStatus,
)
from app.services.integrity import audit, is_commission_code

# Names contain commas ("Regalias Digitales, LLC") and even semicolons, so lists
# of names are separated by a bar.
NAME_SEP = " | "

ROSTER_COLUMNS = [
    ("Entry ID", 9),
    ("Name", 34),
    ("Type", 18),
    ("Payee", 28),
    ("Writer codes", 22),
    ("Commission codes", 22),
    ("Alternate names", 40),
    ("Catalogs", 14),
    ("Cadence", 12),
    ("Language", 9),
    ("Contacts", 40),
    ("Signed in", 30),
    ("Issues", 60),
    ("Action", 30),
    ("Notes", 30),
]
COL = {name: i for i, (name, _) in enumerate(ROSTER_COLUMNS)}

HOW_TO = [
    "How to read and correct this file",
    "",
    "Roster: one row per client or commission partner in the portal. Rows with problems come first, highlighted.",
    "Writer codes (C…, JN…) are a client's own catalog. Commission codes (CS…, CSJ…) belong on a separate "
    "commission-partner entry, even for the same person.",
    "",
    "To correct, edit the Roster sheet and upload it on the Clients page (Upload corrections). You see every change "
    "before anything is applied.",
    "  Move a code: list it on the row it belongs to (and take it off the old row). Its statements move with it.",
    "  New entry: add a row with no Entry ID, a Name and a Type (Client or Commission partner), and list its codes. "
    "Example: the commission side of B-Legit is a new row 'B-Legit', type Commission partner, commission codes "
    "CS0065, CSJ042.",
    "  Alternate names: the names statements may arrive under, separated by ' | '. Delete any that belong to "
    "someone else.",
    "  Name, Type, Payee, Catalogs (MECH, YT, PERF), Cadence (semiannual / quarterly), Language (EN / ES): edit in place.",
    "  Contacts: add an email to give that person access to invite. Removing an email here does not remove access.",
    "  Nothing is ever deleted by an upload. A code listed on no row stays where it is.",
    "",
    "Issues sheet: every problem, one per line.",
    "  MIXED_ROLES — writer and commission codes on one entry; the commission codes need their own entry.",
    "  KIND_MISMATCH — the entry's type does not match its codes.",
    "  ROW_ABSORBED — a client from the client list was filed as another entry's alternate name.",
    "  ALIAS_OF_OTHER — an alternate name on this entry is another client's name.",
    "  ALIAS_UNRELATED — an alternate name shares nothing with the name or payee; check it.",
    "  ROW_MISSING — on the client list but not set up in the portal (left for review at import, or no "
    "statements yet).",
]

_ISSUE_FILL = PatternFill("solid", fgColor="FDECEA")
_HEAD_FILL = PatternFill("solid", fgColor="E8EAED")


def kind_label(kind) -> str:
    if kind == WriterKind.COMMISSION_PARTNER:
        return "Commission partner"
    if kind == WriterKind.CLIENT:
        return "Client"
    return "Unmatched"


def build(db: Session, client_rows=None) -> bytes:
    writers = (
        db.query(Writer)
        .filter(Writer.status != WriterStatus.OFFBOARDED, Writer.kind.isnot(None))
        .all()
    )
    codes = defaultdict(list)
    for a in db.query(BeneficiaryAccount).all():
        codes[a.writer_id].append(a.account_code)
    aliases = defaultdict(list)
    for al in db.query(WriterAlias).all():
        aliases[al.writer_id].append(al.alias_name)
    contacts = defaultdict(list)
    signed_in = defaultdict(list)
    for link, contact in db.query(WriterContact, Contact).join(Contact, Contact.id == WriterContact.contact_id):
        contacts[link.writer_id].append(contact.email)
        if link.user_id is not None:
            signed_in[link.writer_id].append(contact.email)

    findings = audit(db, client_rows=client_rows)
    issues = defaultdict(list)
    for f in findings:
        if f.writer_id is not None:
            issues[f.writer_id].append(f"{f.check}: {f.detail}")

    rows = []
    for w in writers:
        mine = sorted(codes.get(w.id, []))
        rows.append([
            w.id,
            w.canonical_name,
            kind_label(w.kind),
            w.payee_name or "",
            ", ".join(c for c in mine if not is_commission_code(c)),
            ", ".join(c for c in mine if is_commission_code(c)),
            NAME_SEP.join(sorted(a for a in set(aliases.get(w.id, [])) if a != w.canonical_name)),
            ", ".join(w.expected_catalogs or []),
            w.cadence.value if w.cadence is not None else "",
            (w.preferred_language or "").upper(),
            ", ".join(sorted(contacts.get(w.id, []))),
            ", ".join(sorted(signed_in.get(w.id, []))),
            "\n".join(issues.get(w.id, [])),
            "",
            "",
        ])
    # Problems first, then alphabetical: the review starts at the top.
    rows.sort(key=lambda r: (0 if r[COL["Issues"]] else 1, r[COL["Name"]].lower()))

    wb = Workbook()
    ws = wb.active
    ws.title = "Roster"
    ws.append([c for c, _ in ROSTER_COLUMNS])
    for i, (_, width) in enumerate(ROSTER_COLUMNS, start=1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width
        ws.cell(row=1, column=i).font = Font(bold=True)
        ws.cell(row=1, column=i).fill = _HEAD_FILL
    for r in rows:
        ws.append(r)
        if r[COL["Issues"]]:
            for cell in ws[ws.max_row]:
                cell.fill = _ISSUE_FILL
        ws.cell(row=ws.max_row, column=COL["Issues"] + 1).alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions

    ws2 = wb.create_sheet("Issues")
    ws2.append(["Severity", "Check", "Entry ID", "Entry", "Detail"])
    for i, width in enumerate([10, 20, 9, 34, 100], start=1):
        ws2.column_dimensions[ws2.cell(row=1, column=i).column_letter].width = width
        ws2.cell(row=1, column=i).font = Font(bold=True)
        ws2.cell(row=1, column=i).fill = _HEAD_FILL
    for f in findings:
        ws2.append([f.severity, f.check, f.writer_id or "", f.writer, f.detail])
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = ws2.dimensions

    ws3 = wb.create_sheet("How to use")
    for line in HOW_TO:
        ws3.append([line])
    ws3.column_dimensions["A"].width = 130
    ws3["A1"].font = Font(bold=True, size=13)

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def filename() -> str:
    return f"rd-roster-{date.today().isoformat()}.xlsx"
