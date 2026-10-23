"""Roster integrity checks (app/services/integrity.py).

Each case is a shape that reached RD in October 2026: a person's client and
commission codes merged on one entry, and a separate client filed as an alias
of an unrelated one.
"""
from types import SimpleNamespace

import pytest

from app.models.statements import (
    BeneficiaryAccount,
    Catalog,
    Publisher,
    Writer,
    WriterAlias,
    WriterKind,
    WriterStatus,
)
from app.services.client_import.parser import WriterKind as ParsedKind
from app.services.integrity import audit, is_commission_code


@pytest.fixture()
def roster(session):
    pub = Publisher(name="Regalias Digitales")
    session.add(pub)
    session.flush()

    def writer(name, kind=WriterKind.CLIENT, codes=(), payee=None, aliases=(), **kw):
        w = Writer(publisher_id=pub.id, canonical_name=name, kind=kind, payee_name=payee, **kw)
        session.add(w)
        session.flush()
        for code in codes:
            session.add(BeneficiaryAccount(writer_id=w.id, account_code=code, catalog=Catalog.YT))
        for a in aliases:
            session.add(WriterAlias(writer_id=w.id, alias_name=a))
        session.flush()
        return w

    return writer


def checks(findings, name=None):
    return sorted(f.check for f in findings if name is None or f.writer == name)


def row(name, kind, n=1):
    return SimpleNamespace(name=name, kind=kind, row_no=n)


def test_commission_codes_start_with_cs():
    assert is_commission_code("CS0065") and is_commission_code("csj042")
    assert not is_commission_code("C00354") and not is_commission_code("JN0188")


def test_clean_roster_has_no_findings(session, roster):
    roster("B-Legit", codes=["C00354", "JN0188"], payee="Block Movement Media")
    roster("B-Legit", kind=WriterKind.COMMISSION_PARTNER, codes=["CS0065", "CSJ042"], payee="Brandt Jones")
    rows = [row("B-Legit", ParsedKind.CLIENT), row("B-Legit", ParsedKind.COMMISSION_PARTNER)]
    assert audit(session, client_rows=rows) == []


def test_client_and_commission_codes_on_one_entry_is_an_error(session, roster):
    roster("B-Legit", kind=WriterKind.COMMISSION_PARTNER, codes=["C00354", "CS0065", "CSJ042", "JN0188"])
    [f] = audit(session)
    assert (f.check, f.severity) == ("MIXED_ROLES", "error")
    assert "C00354, JN0188" in f.detail and "CS0065, CSJ042" in f.detail


def test_kind_must_match_the_codes(session, roster):
    roster("T.K.", kind=WriterKind.CLIENT, codes=["CS0063", "CSJ039"])
    roster("Partner", kind=WriterKind.COMMISSION_PARTNER, codes=["C00001"])
    assert checks(audit(session)) == ["KIND_MISMATCH", "KIND_MISMATCH"]


def test_client_filed_as_an_alias_of_another_is_caught(session, roster):
    roster("Akwid / AkwidAfterVydia", codes=["C00200"], payee="Akwid", aliases=["Rata Blanca (Adrian Barilari)"])
    roster("Rata Blanca", codes=["C00122", "JN0228"], payee="Rata Blanca")
    rows = [
        row("Akwid / AkwidAfterVydia", ParsedKind.CLIENT),
        row("Rata Blanca", ParsedKind.CLIENT),
        row("Rata Blanca (Adrian Barilari)", ParsedKind.CLIENT),
    ]
    found = audit(session, client_rows=rows)
    assert checks(found, "Akwid / AkwidAfterVydia") == ["ALIAS_UNRELATED", "ROW_ABSORBED"]
    absorbed = next(f for f in found if f.check == "ROW_ABSORBED")
    assert "Rata Blanca (Adrian Barilari)" in absorbed.detail


def test_alias_naming_another_entry_is_an_error(session, roster):
    roster("AmpLive", codes=["C00010"], aliases=["Miguel Mateos"])
    roster("Miguel Mateos", codes=["C00011"])
    [f] = audit(session)
    assert (f.check, f.writer) == ("ALIAS_OF_OTHER", "AmpLive")


def test_payee_and_spelling_aliases_are_fine(session, roster):
    roster("Jay R", codes=["C00676"], payee="Gaudencio Sillona", aliases=["Gaudencio Sillona", "Jay R."])
    assert audit(session) == []


def test_list_row_without_an_entry_is_reported(session, roster):
    roster("Likybo", codes=["C00434"])
    rows = [row("Likybo", ParsedKind.CLIENT), row("Likybo", ParsedKind.COMMISSION_PARTNER, n=7)]
    [f] = audit(session, client_rows=rows)
    assert f.check == "ROW_MISSING" and "commission partner" in f.detail and "row 7" in f.detail


def test_offboarded_entries_are_ignored(session, roster):
    roster("Old", codes=["C00999", "CS0999"], status=WriterStatus.OFFBOARDED)
    assert audit(session) == []


def test_roster_export_puts_problem_entries_first(session, roster):
    import io

    import openpyxl

    from app.services import roster_export

    roster("Aaron", codes=["C00001"])
    roster("B-Legit", kind=WriterKind.COMMISSION_PARTNER, codes=["C00354", "CS0065"])
    wb = openpyxl.load_workbook(io.BytesIO(roster_export.build(session)))
    assert wb.sheetnames == ["Roster", "Issues", "How to use"]
    rows = list(wb["Roster"].iter_rows(min_row=2, values_only=True))
    col = roster_export.COL
    first = rows[0]
    assert first[col["Name"]] == "B-Legit" and first[col["Writer codes"]] == "C00354"
    assert first[col["Commission codes"]] == "CS0065"
    assert first[col["Issues"]].startswith("MIXED_ROLES")
    assert rows[1][col["Name"]] == "Aaron" and not rows[1][col["Issues"]]
    assert [r[1] for r in wb["Issues"].iter_rows(min_row=2, values_only=True)] == ["MIXED_ROLES"]
