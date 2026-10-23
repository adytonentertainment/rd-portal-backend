"""Bulk roster corrections: the export, edited and uploaded back.

The cases are RD's own (October 2026): split B-Legit's commission codes onto a
commission entry, give Rata Blanca (Adrian Barilari) its own entry and take its
name off Akwid — and refuse to repeat the mistake that caused them.
"""
import io

import openpyxl
import pytest

from app.models.statements import (
    BeneficiaryAccount,
    Catalog,
    Contact,
    Publisher,
    Statement,
    Writer,
    WriterAlias,
    WriterContact,
    WriterKind,
)
from app.services import roster_corrections as rc
from app.services import roster_export
from app.services.integrity import audit


@pytest.fixture()
def roster(session):
    pub = Publisher(name="Regalias Digitales")
    session.add(pub)
    session.flush()

    def writer(name, kind=WriterKind.CLIENT, codes=(), aliases=(), statements=0, **kw):
        w = Writer(publisher_id=pub.id, canonical_name=name, kind=kind, **kw)
        session.add(w)
        session.flush()
        for code in codes:
            a = BeneficiaryAccount(writer_id=w.id, account_code=code, catalog=Catalog.YT)
            session.add(a)
            session.flush()
            for i in range(statements):
                session.add(Statement(batch_id=1, account_id=a.id, period_code="PUB26H1", version=i + 1))
        for al in aliases:
            session.add(WriterAlias(writer_id=w.id, alias_name=al))
        session.commit()
        return w

    return writer


def exported(session):
    return openpyxl.load_workbook(io.BytesIO(roster_export.build(session)))


def save(wb) -> bytes:
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def rows_by_name(ws):
    cols = {c.value: i for i, c in enumerate(ws[1])}
    return cols, {r[cols["Name"]].value: r for r in ws.iter_rows(min_row=2)}


def test_unchanged_export_is_a_no_op(session, roster):
    roster("B-Legit", kind=WriterKind.COMMISSION_PARTNER, codes=["C00354", "CS0065"], aliases=["Brandt Jones"])
    roster("Rata Blanca", codes=["C00122", "JN0228"])
    p = rc.plan(session, save(exported(session)))
    assert p.problems == [] and p.ops == []


def test_split_commission_codes_onto_a_new_commission_entry(session, roster):
    b = roster("B-Legit", kind=WriterKind.COMMISSION_PARTNER, codes=["C00354", "JN0188", "CS0065", "CSJ042"],
               statements=2)
    wb = exported(session)
    ws = wb["Roster"]
    cols, rows = rows_by_name(ws)
    row = rows["B-Legit"]
    row[cols["Type"]].value = "Client"
    row[cols["Commission codes"]].value = ""
    ws.append({cols["Name"] + 1: "B-Legit", cols["Type"] + 1: "Commission partner",
               cols["Commission codes"] + 1: "CS0065, CSJ042", cols["Contacts"] + 1: "blegit72@gmail.com"})
    data = save(wb)

    p = rc.plan(session, data)
    assert p.problems == []
    texts = [c["text"] for c in p.summary()["changes"]]
    assert any(t.startswith("Create new commission partner “B-Legit”") for t in texts)
    assert any("Change #%d B-Legit from Commission partner to Client" % b.id in t for t in texts)
    assert any(t.startswith("Move CS0065 (2 statements)") for t in texts)

    result = rc.apply(session, data)
    [new_id] = result["created_ids"]
    codes = {a.account_code: a.writer_id for a in session.query(BeneficiaryAccount)}
    assert codes["CS0065"] == codes["CSJ042"] == new_id
    assert codes["C00354"] == codes["JN0188"] == b.id
    assert session.get(Writer, b.id).kind == WriterKind.CLIENT
    assert session.get(Writer, new_id).kind == WriterKind.COMMISSION_PARTNER
    contact = session.query(Contact).filter_by(email="blegit72@gmail.com").one()
    assert session.query(WriterContact).filter_by(writer_id=new_id, contact_id=contact.id).count() == 1
    assert [f.check for f in audit(session)] == []
    # Uploading the same file again changes nothing more... except it would try
    # to create the commission row again — so an export taken now is the base.
    assert rc.plan(session, save(exported(session))).ops == []


def test_absorbed_client_gets_its_own_entry_and_the_alias_comes_off(session, roster):
    akwid = roster("Akwid", codes=["C00200"], aliases=["Rata Blanca (Adrian Barilari)", "Akwid"])
    rata = roster("Rata Blanca", codes=["C00122", "JN0228"])
    wb = exported(session)
    ws = wb["Roster"]
    cols, rows = rows_by_name(ws)
    rows["Akwid"][cols["Alternate names"]].value = ""
    rows["Rata Blanca"][cols["Writer codes"]].value = "JN0228"
    ws.append({cols["Name"] + 1: "Rata Blanca (Adrian Barilari)", cols["Type"] + 1: "Client",
               cols["Writer codes"] + 1: "C00122", cols["Language"] + 1: "ES"})
    rc.apply(session, save(wb))
    barilari = session.query(Writer).filter_by(canonical_name="Rata Blanca (Adrian Barilari)").one()
    assert barilari.preferred_language == "es"
    assert session.query(BeneficiaryAccount).filter_by(account_code="C00122").one().writer_id == barilari.id
    assert session.query(BeneficiaryAccount).filter_by(account_code="JN0228").one().writer_id == rata.id
    assert [a.alias_name for a in session.query(WriterAlias).filter_by(writer_id=akwid.id)] == ["Akwid"]


def test_the_original_mistake_is_refused(session, roster):
    roster("B-Legit", codes=["C00354"])
    wb = exported(session)
    ws = wb["Roster"]
    cols, rows = rows_by_name(ws)
    rows["B-Legit"][cols["Commission codes"]].value = "CS0065"
    p = rc.plan(session, save(wb))
    assert any("commission code on a client" in x for x in p.problems)
    with pytest.raises(rc.CorrectionsRefused):
        rc.apply(session, save(wb))
    assert session.query(BeneficiaryAccount).count() == 1


def test_code_in_the_wrong_column_or_on_two_rows_is_refused(session, roster):
    roster("A", codes=["C00001"])
    roster("B", codes=["C00002"])
    wb = exported(session)
    ws = wb["Roster"]
    cols, rows = rows_by_name(ws)
    rows["B"][cols["Writer codes"]].value = "C00002, C00001"
    rows["A"][cols["Writer codes"]].value = "C00001, CS0001"
    problems = rc.plan(session, save(wb)).problems
    assert any("C00001 is listed on rows" in x for x in problems)
    assert any("CS0001 is a commission code; put it under Commission codes" in x for x in problems)


def test_a_code_taken_off_a_row_and_listed_nowhere_stays_with_a_warning(session, roster):
    roster("A", codes=["C00001", "C00002"])
    wb = exported(session)
    ws = wb["Roster"]
    cols, rows = rows_by_name(ws)
    rows["A"][cols["Writer codes"]].value = "C00001"
    p = rc.plan(session, save(wb))
    assert p.ops == [] and p.problems == []
    assert any("keeps C00002" in w for w in p.warnings)


def test_not_a_roster_file(session, roster):
    wb = openpyxl.Workbook()
    wb.active.append(["Artist", "Email"])
    assert "missing column" in rc.plan(session, save(wb)).problems[0]
    assert "not a spreadsheet" in rc.plan(session, b"hello").problems[0]


def test_endpoints_preview_then_apply(session, roster):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from app.database.session import get_session
    from app.models.models import User
    from app.routers.auth import get_user
    from app.routers.writers_admin import writers_admin_router

    roster("Rata Blanca", codes=["C00122", "JN0228"])
    admin = User(email="a@example.com", username="a", role="admin", admin_approved=True, royalty_per_stream=0)
    session.add(admin)
    session.commit()
    app = FastAPI()
    app.include_router(writers_admin_router)
    app.dependency_overrides[get_session] = lambda: session
    app.dependency_overrides[get_user] = lambda: admin
    api = TestClient(app)

    wb = exported(session)
    ws = wb["Roster"]
    cols, rows = rows_by_name(ws)
    rows["Rata Blanca"][cols["Writer codes"]].value = "JN0228"
    ws.append({cols["Name"] + 1: "Rata Blanca (Adrian Barilari)", cols["Type"] + 1: "Client",
               cols["Writer codes"] + 1: "C00122"})
    files = {"file": ("roster.xlsx", save(wb), "application/octet-stream")}

    preview = api.post("/admin/writers/corrections/preview", files=files).json()
    assert preview["problems"] == [] and preview["counts"] == {"create": 1, "move": 1}
    assert session.query(Writer).count() == 1  # preview changes nothing

    applied = api.post("/admin/writers/corrections/apply", files=files).json()
    assert applied["applied"] == 2 and applied["remaining_errors"] == 0

    bad = exported(session)
    cols, rows = rows_by_name(bad["Roster"])
    rows["Rata Blanca"][cols["Commission codes"]].value = "CS0001"
    r = api.post("/admin/writers/corrections/apply", files={"file": ("x.xlsx", save(bad), "application/octet-stream")})
    assert r.status_code == 422 and r.json()["detail"].startswith("Nothing was changed")
